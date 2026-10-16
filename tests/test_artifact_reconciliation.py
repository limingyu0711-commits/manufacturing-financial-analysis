from __future__ import annotations

import json
import unittest
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook


PROJECT_ROOT = Path(__file__).resolve().parents[1]
DATA_DIR = PROJECT_ROOT / "data"


class ExportedArtifactReconciliationTest(unittest.TestCase):
    """Validate the files a reviewer will actually open, not only in-memory data."""

    @classmethod
    def setUpClass(cls) -> None:
        cls.csv = pd.read_csv(DATA_DIR / "manufacturing_financial_data.csv")
        cls.dashboard = json.loads(
            (DATA_DIR / "dashboard_data.json").read_text(encoding="utf-8")
        )
        cls.workbook = load_workbook(
            DATA_DIR / "manufacturing_financial_analysis.xlsx",
            data_only=False,
        )
        cls.workbook_values = load_workbook(
            DATA_DIR / "manufacturing_financial_analysis.xlsx",
            data_only=True,
        )

    def test_csv_has_complete_unique_controlled_grain(self) -> None:
        self.assertEqual(len(self.csv), 36)
        self.assertEqual(self.csv["Month"].nunique(), 12)
        self.assertEqual(self.csv["Product Line"].nunique(), 3)
        self.assertFalse(self.csv.duplicated(["Month", "Product Line"]).any())

    def test_csv_financial_identities_hold_after_export_rounding(self) -> None:
        for scenario in ("Budget", "Actual"):
            revenue_difference = (
                self.csv[f"{scenario} Revenue"]
                - self.csv[f"{scenario} Sales Volume"]
                * self.csv[f"{scenario} Unit Price"]
            ).abs()
            cogs_difference = (
                self.csv[f"{scenario} COGS"]
                - self.csv[f"{scenario} Material Cost"]
                - self.csv[f"{scenario} Labor Cost"]
                - self.csv[f"{scenario} Manufacturing Overhead"]
            ).abs()
            gross_profit_difference = (
                self.csv[f"{scenario} Gross Profit"]
                - self.csv[f"{scenario} Revenue"]
                + self.csv[f"{scenario} COGS"]
            ).abs()
            self.assertLessEqual(revenue_difference.max(), 0.01)
            self.assertLessEqual(cogs_difference.max(), 0.01)
            self.assertLessEqual(gross_profit_difference.max(), 0.01)

            expected_margin = (
                self.csv[f"{scenario} Gross Profit"]
                / self.csv[f"{scenario} Revenue"]
            )
            self.assertTrue(
                np.allclose(
                    self.csv[f"{scenario} Gross Margin"],
                    expected_margin,
                    atol=1e-6,
                )
            )

    def test_month_product_and_annual_totals_reconcile(self) -> None:
        annual_revenue = self.csv["Actual Revenue"].sum()
        annual_cogs = self.csv["Actual COGS"].sum()
        self.assertTrue(
            np.isclose(
                self.csv.groupby("Month")["Actual Revenue"].sum().sum(),
                annual_revenue,
            )
        )
        self.assertTrue(
            np.isclose(
                self.csv.groupby("Product Line")["Actual Revenue"].sum().sum(),
                annual_revenue,
            )
        )
        self.assertTrue(
            np.isclose(
                (
                    self.csv["Actual Material Cost"]
                    + self.csv["Actual Labor Cost"]
                    + self.csv["Actual Manufacturing Overhead"]
                ).sum(),
                annual_cogs,
            )
        )

    def test_dashboard_json_matches_exported_csv(self) -> None:
        metadata = self.dashboard["metadata"]
        self.assertEqual(metadata["dataClassification"], "Synthetic Data")
        self.assertEqual(metadata["companyType"], "Simulated Company")

        json_frame = pd.DataFrame(self.dashboard["records"]).sort_values(
            ["Month", "Product Line"]
        )
        csv_frame = self.csv.sort_values(["Month", "Product Line"])
        self.assertEqual(len(json_frame), len(csv_frame))
        for column in (
            "Budget Revenue",
            "Actual Revenue",
            "Budget COGS",
            "Actual COGS",
            "Budget Gross Profit",
            "Actual Gross Profit",
            "Gross Margin Variance (pp)",
        ):
            self.assertTrue(
                np.allclose(json_frame[column], csv_frame[column], atol=1e-9)
            )

    def test_workbook_audit_controls_pass(self) -> None:
        audit = self.workbook_values["Audit"]
        statuses = [audit.cell(row=row, column=4).value for row in range(6, 13)]
        self.assertEqual(statuses, ["PASS"] * 7)

    def test_workbook_summary_reconciles_to_csv(self) -> None:
        summary = self.workbook_values["Summary"]
        expected = {
            "Revenue": (
                self.csv["Budget Revenue"].sum(),
                self.csv["Actual Revenue"].sum(),
            ),
            "COGS": (
                self.csv["Budget COGS"].sum(),
                self.csv["Actual COGS"].sum(),
            ),
            "Gross Profit": (
                self.csv["Budget Gross Profit"].sum(),
                self.csv["Actual Gross Profit"].sum(),
            ),
        }
        for row in range(6, 9):
            metric = summary.cell(row=row, column=1).value
            budget = summary.cell(row=row, column=2).value
            actual = summary.cell(row=row, column=3).value
            self.assertTrue(np.isclose(budget, expected[metric][0], atol=0.01))
            self.assertTrue(np.isclose(actual, expected[metric][1], atol=0.01))

        budget_margin = self.csv["Budget Gross Profit"].sum() / self.csv[
            "Budget Revenue"
        ].sum()
        actual_margin = self.csv["Actual Gross Profit"].sum() / self.csv[
            "Actual Revenue"
        ].sum()
        self.assertTrue(np.isclose(summary["B9"].value, budget_margin, atol=1e-9))
        self.assertTrue(np.isclose(summary["C9"].value, actual_margin, atol=1e-9))

    def test_workbook_is_formula_driven(self) -> None:
        data = self.workbook["Data"]
        summary = self.workbook["Summary"]
        self.assertTrue(str(data["P6"].value).startswith("="))
        self.assertTrue(str(data["AD6"].value).startswith("="))
        self.assertTrue(str(summary["B6"].value).startswith("="))
        self.assertTrue(str(summary["B9"].value).startswith("="))

    def test_exported_summary_keeps_margin_precision(self) -> None:
        summary = pd.read_csv(PROJECT_ROOT / "reports" / "financial_summary.csv")
        row = summary.loc[summary["Metric"].eq("Gross Margin")].iloc[0]
        expected_budget = self.csv["Budget Gross Profit"].sum() / self.csv[
            "Budget Revenue"
        ].sum()
        expected_actual = self.csv["Actual Gross Profit"].sum() / self.csv[
            "Actual Revenue"
        ].sum()
        self.assertTrue(np.isclose(row["Budget"], expected_budget, atol=5e-7))
        self.assertTrue(np.isclose(row["Actual"], expected_actual, atol=5e-7))


if __name__ == "__main__":
    unittest.main()
